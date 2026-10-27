"""Excel ফাইল যাচাই, সংরক্ষণ ও নিরাপদ মোছা।"""
import os
import shutil
import uuid
from pathlib import Path

from openpyxl import load_workbook

import config
import db
import history

ACTIVE_QUEUE = ("pending", "processing", "uncertain")


def read_headers_and_rows(path: Path):
    """(কলামের নাম তালিকা, ডাটা সারি সংখ্যা) ফেরত দেয়। পড়া না গেলে Exception।"""
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        headers = [str(c).strip() for c in (header or ()) if c is not None and str(c).strip()]
        count = sum(1 for r in rows if any(c is not None and str(c).strip() for c in r))
        return headers, count
    finally:
        wb.close()


def parse_columns(text: str) -> list[str]:
    return [c.strip() for c in (text or "").replace("،", ",").split(",") if c.strip()]


def validate(path: Path, required_columns: list[str] | None = None):
    """(errors, headers, row_count) — errors খালি হলে ফাইল ঠিক আছে।"""
    errors: list[str] = []
    if path.suffix.lower() != ".xlsx":
        return ["ফাইলটি .xlsx ফরম্যাটে হতে হবে।"], [], 0
    if path.stat().st_size > config.MAX_EXCEL_BYTES:
        return [f"ফাইল অনেক বড় (সর্বোচ্চ {config.MAX_EXCEL_BYTES // (1024 * 1024)} MB)।"], [], 0
    try:
        headers, count = read_headers_and_rows(path)
    except Exception:
        return ["Excel ফাইলটি পড়া যায়নি। ফাইল নষ্ট বা ভুল ফরম্যাটের হতে পারে।"], [], 0
    if not headers:
        errors.append("প্রথম সারিতে কোনো কলামের নাম (Header) পাওয়া যায়নি।")
    if count == 0:
        errors.append("ফাইলে কোনো ডাটা সারি নেই।")
    if required_columns and headers:
        have = {h.lower() for h in headers}
        missing = [c for c in required_columns if c.lower() not in have]
        if missing:
            errors.append("এই প্রয়োজনীয় কলামগুলো নেই: " + ", ".join(missing))
    return errors, headers, count


def new_tmp_path() -> Path:
    return config.TMP_DIR / f"{uuid.uuid4().hex}.xlsx"


def discard(path) -> None:
    try:
        if path and Path(path).exists():
            Path(path).unlink()
    except OSError:
        pass


def move_to_storage(tmp_path: Path) -> Path:
    """অস্থায়ী ফাইল স্থায়ী ফোল্ডারে নেয় (ইউনিক নামে)।"""
    final = config.EXCEL_DIR / f"{uuid.uuid4().hex}.xlsx"
    shutil.move(str(tmp_path), str(final))
    return final


def get_active(profile_id: int):
    return db.one("SELECT * FROM excel_files WHERE profile_id=? AND status='active' "
                  "ORDER BY id DESC LIMIT 1", (profile_id,))


def in_use(excel_id: int, exclude_queue_id: int | None = None) -> bool:
    q = ("SELECT 1 FROM job_queue WHERE excel_file_id=? AND status IN "
         "('pending','processing','uncertain')")
    params = [excel_id]
    if exclude_queue_id:
        q += " AND id<>?"
        params.append(exclude_queue_id)
    return db.one(q, tuple(params)) is not None


def register(conn, profile_id: int, stored_path: Path, original_name: str, rows: int) -> int:
    """নতুন Excel রেকর্ড (কলার ট্রানজ্যাকশন চালাবে)। আগের active ফাইল 'replaced' হয়।"""
    prev = conn.execute("SELECT id, file_path FROM excel_files WHERE profile_id=? AND status='active'",
                        (profile_id,)).fetchall()
    for p in prev:
        conn.execute("UPDATE excel_files SET status='replaced', deleted_at=? WHERE id=?",
                     (db.now_utc(), p["id"]))
    cur = conn.execute(
        "INSERT INTO excel_files(profile_id, file_path, original_name, row_count, status, created_at) "
        "VALUES(?,?,?,?, 'active', ?)", (profile_id, str(stored_path), original_name, rows, db.now_utc()))
    return cur.lastrowid


def replace_excel(profile_id: int, tmp_path: Path, original_name: str, rows: int):
    """বিদ্যমান Profile-এ Excel বদলায়। Queue-তে কাজ থাকলে বদলাবে না।
    ফেরত: (ok, বার্তা)"""
    cur = get_active(profile_id)
    if cur and in_use(cur["id"]):
        return False, "এই Excel-এর সাথে Queue-তে অসম্পূর্ণ কাজ আছে। আগে সেগুলো সম্পন্ন বা বাতিল করুন।"
    stored = move_to_storage(tmp_path)
    try:
        with db.tx() as conn:
            register(conn, profile_id, stored, original_name, rows)
    except Exception:
        discard(stored)
        raise
    if cur:
        discard(cur["file_path"])
    history.add("excel_uploaded", f"Excel আপলোড: {original_name} ({rows} সারি)", profile_id)
    return True, "Excel সংরক্ষিত হয়েছে।"


def delete_after_success(item) -> tuple[bool, str]:
    """সফল Submission-এর পর Excel মোছে — তবে আগে মিলিয়ে দেখে:
    ফাইলটি এই Queue Item ও Profile-এর, active, এবং অন্য কোনো অসম্পূর্ণ কাজ এটি ব্যবহার করছে না।"""
    eid = item["excel_file_id"]
    if not eid:
        return False, "Queue Item-এ কোনো Excel যুক্ত ছিল না।"
    row = db.one("SELECT * FROM excel_files WHERE id=?", (eid,))
    if not row or row["profile_id"] != item["profile_id"]:
        return False, "Excel ও Profile মেলেনি — ফাইল মোছা হয়নি।"
    if row["status"] != "active":
        return False, "Excel আগেই মোছা/বদলানো হয়েছে।"
    if in_use(eid, exclude_queue_id=item["id"]):
        return False, "অন্য Job-এর কাজ এই Excel ব্যবহার করছে — ফাইল রাখা হয়েছে।"
    try:
        if os.path.exists(row["file_path"]):
            os.remove(row["file_path"])
    except OSError as e:
        history.add("error", f"Excel মোছা যায়নি: {e}", item["profile_id"], item["job_id"])
        return False, "Excel ফাইল মোছা যায়নি।"
    db.run("UPDATE excel_files SET status='deleted', deleted_at=? WHERE id=?", (db.now_utc(), eid))
    history.add("excel_deleted", f"সফল Submission-এর পর Excel মোছা হয়েছে ({row['original_name']})",
                item["profile_id"], item["job_id"])
    return True, "Excel মোছা হয়েছে।"
